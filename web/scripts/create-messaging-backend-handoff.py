"""Build the editable RTL backend handoff from its single JSON source.
Run using the bundled Python runtime; PDF export uses the companion Word script.
"""
from pathlib import Path
import json, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / 'docs/product/messaging-backend-policy-20261101.json'
STEM = 'MakanApp_Backend_Messaging_Policy_FA_20261101'
FONT = 'Tahoma'

def prop(parent, tag, val=None):
    old = parent.find(qn(tag))
    element = old if old is not None else OxmlElement(tag)
    if old is None: parent.append(element)
    if val is not None: element.set(qn('w:val'), str(val))
    return element

def rtl(p, enabled=True):
    prop(p._p.get_or_add_pPr(), 'w:bidi', '1' if enabled else '0')
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if enabled else WD_ALIGN_PARAGRAPH.LEFT

def text(p, value, size=11, bold=False):
    # Separate Latin API identifiers from Persian runs for stable mixed-direction text.
    for part in re.split(r'([A-Za-z][A-Za-z0-9_./{}=:+?@-]*(?: [A-Za-z0-9_./{}=:+?@-]+)*)', value):
        if not part: continue
        latin = bool(re.match(r'^[A-Za-z]', part))
        r = p.add_run(part); r.bold = bold; r.font.name = 'Consolas' if latin else FONT
        r.font.size = Pt(size-0.7 if latin else size); r.font.color.rgb = RGBColor(0,0,0)
        pr = r._r.get_or_add_rPr(); fonts = prop(pr,'w:rFonts')
        for a in ['ascii','hAnsi','cs','eastAsia']: fonts.set(qn('w:'+a), 'Consolas' if latin else FONT)
        prop(pr,'w:rtl','0' if latin else '1');prop(pr,'w:szCs',round(size*2));prop(pr,'w:lang','fa-IR' if not latin else 'en-US')
    return p

def paragraph(doc, value, style=None, size=11):
    p = doc.add_paragraph(style=style);rtl(p);text(p,value,size)
    p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1.3
    return p

def table(doc, block):
    headers, rows = block['headers'], block['rows']
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.RIGHT;t.autofit=False
    tp=t._tbl.tblPr;prop(tp,'w:bidiVisual','1')
    borders=prop(tp,'w:tblBorders')
    for tag in ['top','left','bottom','right','insideH','insideV']:
        edge=prop(borders,'w:'+tag,'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9')
    margin=prop(tp,'w:tblCellMar')
    for tag in ['top','bottom','left','right']:
        item=prop(margin,'w:'+tag);item.set(qn('w:w'),'40' if headers[0]=='شماره' and tag in ['top','bottom'] else '75');item.set(qn('w:type'),'dxa')
    n=len(headers)
    if headers[0]=='شماره': widths=[.52,2.08,4.60]
    elif headers[0]=='شناسه':widths=[.55,2.20,4.45]
    elif n==4:widths=[2.3,1.4,1.4,2.1]
    else:widths=[2.1,2.55,2.55]
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,values in enumerate([headers]+rows):
        row=t.rows[0] if i==0 else t.add_row()
        prop(row._tr.get_or_add_trPr(),'w:cantSplit')
        for j,(cell,value,w) in enumerate(zip(row.cells,values,widths)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade=prop(cell._tc.get_or_add_tcPr(),'w:shd');shade.set(qn('w:fill'),'EAF0FA' if i==0 else ('F8FAFC' if i%2==0 else 'FFFFFF'))
            p=cell.paragraphs[0];rtl(p);p.paragraph_format.space_after=Pt(1 if headers[0]=='شماره' else 2);p.paragraph_format.space_before=Pt(1 if headers[0]=='شماره' else 2);p.paragraph_format.line_spacing=1.12
            text(p,value,9.8 if headers[0]=='شماره' else 10.1,i==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def build():
    source=json.loads(DATA.read_text(encoding='utf-8-sig'));doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.65);sec.right_margin=Inches(.65)
    sec.header_distance=Inches(.27);sec.footer_distance=Inches(.28)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        sty=doc.styles[name];sty.font.name=FONT;sty.font.color.rgb=RGBColor(0,0,0);sty.font.size=Pt(11);sty.font.italic=False
        prop(sty.element.get_or_add_pPr(),'w:bidi','1')
        fonts=prop(sty.element.get_or_add_rPr(),'w:rFonts');fonts.set(qn('w:cs'),FONT)
    for border in doc.styles.element.findall('.//' + qn('w:pBdr')): border.getparent().remove(border)
    for name,size in [('Title',20),('Heading 1',16),('Heading 2',12)]:
        doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True;doc.styles[name].paragraph_format.keep_with_next=True
    hp=sec.header.paragraphs[0];rtl(hp);text(hp,'ماکان   درخواست توسعه برای تیم بک‌اند',8.5)
    fp=sec.footer.paragraphs[0];rtl(fp);fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;text(fp,'ماکان  |  '+source['date']+'  |  صفحه ',8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
    doc.core_properties.title=source['title'];doc.core_properties.author='MakanApp';doc.core_properties.last_modified_by='MakanApp'
    md=['# '+source['title'],'',source['version']+' | '+source['date'],'']
    digits=str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹')
    for index,page in enumerate(source['pages'],1):
        if index>1:doc.add_page_break()
        if index==1:
            paragraph(doc,source['title'],'Title',20)
            paragraph(doc,source['version']+'   '+source['date'],'Subtitle',10)
        paragraph(doc,str(index).translate(digits)+' '+page['title'],'Heading 1',16)
        md.extend(['## '+str(index)+' '+page['title'],''])
        for block in page['blocks']:
            kind=block['type']
            if kind=='p':paragraph(doc,block['text']);md.extend([block['text'],''])
            elif kind=='h':paragraph(doc,block['text'],'Heading 2',12);md.extend(['### '+block['text'],''])
            elif kind=='bullets':
                for item in block['items']:
                    p=paragraph(doc,'• '+item);p.paragraph_format.right_indent=Inches(.08);md.append('- '+item)
                md.append('')
            elif kind=='code':
                p=doc.add_paragraph();rtl(p,False);p.paragraph_format.space_after=Pt(9);p.paragraph_format.line_spacing=1.12
                for line in block['text'].splitlines():
                    r=p.add_run(line+'\n');r.font.name='Consolas';r.font.size=Pt(8)
                    pr=r._r.get_or_add_rPr();prop(pr,'w:rtl','0');prop(pr,'w:cs','0')
                    fonts=prop(pr,'w:rFonts')
                    for a in ['ascii','hAnsi','cs','eastAsia']:fonts.set(qn('w:'+a),'Consolas')
                    lang=prop(pr,'w:lang','en-US');lang.set(qn('w:bidi'),'en-US')
                md.extend(['```text',block['text'],'```',''])
            elif kind=='table':
                table(doc,block);md.append('| '+' | '.join(block['headers'])+' |');md.append('| '+' | '.join(['---']*len(block['headers']))+' |')
                for row in block['rows']:md.append('| '+' | '.join(row)+' |')
                md.append('')
    out=ROOT/'output/docx'/f'{STEM}.docx';out.parent.mkdir(parents=True,exist_ok=True);doc.save(out)
    markdown=ROOT/'docs/product'/'MESSENGER_BACKEND_REQUEST_FA_20261101.md';markdown.write_text('\n'.join(md).rstrip()+'\n',encoding='utf-8')
    print(json.dumps({'docx':str(out),'source':str(DATA),'markdown':str(markdown),'sections':len(source['pages'])}))
if __name__=='__main__': build()
